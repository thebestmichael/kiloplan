from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from kiloplan.api.main import app

client = TestClient(app)

PAUSED_TRIPLE = {
    "pr_weight_kg": 180,
    "pr_reps": 3,
    "pr_rpe": 9,
    "target_reps": 3,
    "target_rpe": 5,
}


def test_health():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_target_load_defaults_to_hybrid():
    response = client.post("/api/target-load", json=PAUSED_TRIPLE)
    assert response.status_code == 200
    body = response.json()
    assert body["e1rm_kg"] == pytest.approx(196.3636, abs=1e-4)
    assert body["target_weight_kg"] == pytest.approx(158.1818, abs=1e-4)


def test_target_load_with_epley():
    response = client.post("/api/target-load", json={**PAUSED_TRIPLE, "formula": "epley"})
    assert response.status_code == 200
    body = response.json()
    assert body["e1rm_kg"] == pytest.approx(204.0)
    assert body["target_weight_kg"] == pytest.approx(161.0526, abs=1e-4)


def test_target_load_with_brzycki():
    response = client.post("/api/target-load", json={**PAUSED_TRIPLE, "formula": "brzycki"})
    assert response.status_code == 200
    body = response.json()
    assert body["target_weight_kg"] == pytest.approx(158.1818, abs=1e-4)


@pytest.mark.parametrize(
    "override",
    [
        {"pr_weight_kg": 0},
        {"pr_reps": 0},
        {"pr_rpe": 10.5},
        {"target_rpe": 0},
        {"formula": "lombardi"},
    ],
)
def test_target_load_rejects_invalid_input(override):
    response = client.post("/api/target-load", json={**PAUSED_TRIPLE, **override})
    assert response.status_code == 422


def test_target_load_rejects_reps_beyond_brzycki_limit():
    response = client.post(
        "/api/target-load", json={**PAUSED_TRIPLE, "formula": "brzycki", "target_reps": 40}
    )
    assert response.status_code == 422
    assert "at most 36" in response.json()["detail"]


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Block A"
    for column, label, dates, load in (
        (2, "Week 1", "7/13-19/26", 298),
        (10, "Week 2", "7/20-26/26", 310),
    ):
        sheet.cell(row=1, column=column, value=label)
        sheet.cell(row=2, column=column, value="Date")
        sheet.cell(row=2, column=column + 1, value=dates)
        header = [
            "Day 1",
            "Sets",
            "Reps",
            "Intensity",
            "Target Load",
            "Actual load",
            "Last Set RPE",
        ]
        line = ["Squats", 1, 7, "@5", None, load, 11 if column == 10 else "<5"]
        for offset, (title, value) in enumerate(zip(header, line, strict=True)):
            sheet.cell(row=4, column=column + offset, value=title)
            sheet.cell(row=5, column=column + offset, value=value)
    workbook.create_sheet("Reference")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_program():
    response = client.post("/api/program/parse", files={"file": ("program.xlsx", workbook_bytes())})
    assert response.status_code == 200
    body = response.json()
    assert body["sheets"] == [
        {
            "name": "Block A",
            "lines": 2,
            "weeks": 2,
            "first_week": "2026-07-13",
            "last_week": "2026-07-20",
        }
    ]
    assert body["skipped_sheets"] == ["Reference"]
    assert [(issue["sheet"], issue["cell"]) for issue in body["issues"]] == [("Block A", "J5")]

    first = body["rows"][0]
    assert (first["movement"], first["sets"], first["reps"]) == ("Squats", 1, 7)
    assert first["week_start"] == "2026-07-13"
    assert first["actual_load_kg"] == pytest.approx(298 * 0.45359237)
    assert (first["rpe"], first["rpe_below_5"]) == (None, True)


def test_parse_program_rejects_non_workbook():
    response = client.post(
        "/api/program/parse", files={"file": ("notes.xlsx", b"this is not a spreadsheet")}
    )
    assert response.status_code == 400
    assert "Excel workbook" in response.json()["detail"]
