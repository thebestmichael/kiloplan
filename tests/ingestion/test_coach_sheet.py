from datetime import date, datetime

import pytest
from openpyxl import Workbook

from kiloplan.ingestion.coach_sheet import (
    KG_PER_LB,
    _parse_rpe,
    _parse_week_dates,
    parse_workbook,
)
from workbooks import NEW_HEADER, add_old_week, put, save


def lb(pounds):
    return pytest.approx(pounds * KG_PER_LB)


@pytest.fixture
def old_layout(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Block A"
    add_old_week(
        sheet,
        2,
        "Week 1",
        "7/13-19/26",
        [
            (
                "Day 1",
                [
                    ["Squats", 1, 7, "@5", None, 298, "<5"],
                    [None, 3, 7, -10, 268.2, 270, 6.5],
                    ["Cable Rows", 4, "AMRAP", "@10", None, "60*", None],
                    ["Notes ", "Low back extension"],
                    ["after the notes", 1, 1, None, None, 999, 9],
                ],
            ),
            ("Day 2", [["Bench", 1, 8, "@3", None, 221, 7]]),
        ],
    )
    add_old_week(
        sheet, 10, "Week 2", "7/20-26/26", [("Day 1", [["Squats", 1, 7, None, None, 310]])]
    )
    workbook.create_sheet("Reference")["B2"] = "Current Maxes"
    return save(workbook, tmp_path)


def test_old_layout_rows(old_layout):
    result = parse_workbook(old_layout)
    assert result.skipped_sheets == ["Reference"]
    assert result.issues == []
    assert [(row.week_label, row.day_label, row.movement) for row in result.rows] == [
        ("Week 1", "Day 1", "Squats"),
        ("Week 1", "Day 1", "Squats"),
        ("Week 1", "Day 1", "Cable Rows"),
        ("Week 1", "Day 2", "Bench"),
        ("Week 2", "Day 1", "Squats"),
    ]

    top_set, backoff, rows, bench, _ = result.rows
    assert (top_set.sheet, top_set.cell) == ("Block A", "B5")
    assert (top_set.week_start, top_set.week_end) == (date(2026, 7, 13), date(2026, 7, 19))
    assert (top_set.sets, top_set.reps, top_set.intensity_raw) == (1, 7, "@5")
    assert top_set.actual_load_kg == lb(298)
    assert (top_set.rpe, top_set.rpe_below_5) == (None, True)

    # a blank movement cell carries the movement down from the row above
    assert (backoff.sets, backoff.reps, backoff.intensity_raw) == (3, 7, "-10")
    assert backoff.target_load_raw == "268.2"
    assert (backoff.rpe, backoff.rpe_below_5) == (6.5, False)

    # text reps stay as text; an annotated load keeps its number
    assert (rows.reps, rows.reps_raw) == (None, "AMRAP")
    assert (rows.actual_load_raw, rows.actual_load_kg) == ("60*", lb(60))
    assert (rows.rpe, rows.rpe_below_5) == (None, False)

    assert bench.week_start == date(2026, 7, 13)
    assert result.rows[-1].week_start == date(2026, 7, 20)


def test_old_layout_weeks_and_notes(old_layout):
    result = parse_workbook(old_layout)
    assert [(week.sheet, week.week_index, week.label, week.start) for week in result.weeks] == [
        ("Block A", 0, "Week 1", date(2026, 7, 13)),
        ("Block A", 1, "Week 2", date(2026, 7, 20)),
    ]
    assert [row.week_index for row in result.rows] == [0, 0, 0, 0, 1]
    # the text beside "Notes" and everything under it, up to the next day
    assert [(note.week_index, note.day_label, note.text) for note in result.notes] == [
        (0, "Day 1", "Low back extension"),
        (0, "Day 1", "after the notes 1 1 999 9"),
    ]


@pytest.fixture
def new_layout(tmp_path):
    """Newest tab layout: weekday names, a Movement column, and no dates filled in."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "5w "
    for column, label in ((2, "Week 1"), (11, "Week 2")):
        put(sheet, 4, column, ["Meet:", "None"])
        put(sheet, 5, column, ["Date:"])
        put(sheet, 7, column, [label])
    put(sheet, 8, 2, ["Monday ", *NEW_HEADER])
    put(sheet, 9, 3, ["Paused Squats", 1, 1, None, "8 / 418 cap", 418, 5.5])
    put(sheet, 10, 3, [None, 3, 4, -0.12, 367.84, 367, 11])
    # Excel turned a rep count of 8 and a target of "8-9" into dates
    put(sheet, 11, 3, ["DB skullcrusher", 4, datetime(1900, 1, 8), "@9", datetime(2026, 8, 9), 63])
    put(sheet, 13, 2, ["Notes"])
    put(sheet, 14, 2, ["420-430w4"])
    # this day has no header row of its own
    put(sheet, 16, 2, ["Saturday", "Bench", 1, 1, None, 8, 331, 6.5])
    put(sheet, 17, 3, [None, 3, 6, -0.05, 251.75, None, 5.5])
    put(sheet, 8, 11, ["Monday ", *NEW_HEADER])
    put(sheet, 9, 12, ["Paused Squats", 1, 2, None, 6, 353, 5])
    return save(workbook, tmp_path)


def test_new_layout_rows(new_layout):
    result = parse_workbook(new_layout)
    assert [(row.week_label, row.day_label, row.movement, row.cell) for row in result.rows] == [
        ("Week 1", "Monday", "Paused Squats", "C9"),
        ("Week 1", "Monday", "Paused Squats", "C10"),
        ("Week 1", "Monday", "DB skullcrusher", "C11"),
        ("Week 1", "Saturday", "Bench", "C16"),
        ("Week 1", "Saturday", "Bench", "C17"),
        ("Week 2", "Monday", "Paused Squats", "L9"),
    ]
    assert all(row.week_start is None for row in result.rows)

    single, backoff, skullcrusher, bench, bench_backoff, _ = result.rows
    assert (single.target_load_raw, single.actual_load_kg, single.rpe) == (
        "8 / 418 cap",
        lb(418),
        5.5,
    )
    assert (skullcrusher.reps, skullcrusher.target_load_raw) == (8, "8-9")
    assert (bench.sets, bench.reps, bench.actual_load_kg, bench.rpe) == (1, 1, lb(331), 6.5)
    assert (bench_backoff.reps, bench_backoff.actual_load_kg, bench_backoff.rpe) == (6, None, 5.5)

    # an RPE outside "<5" and 5 to 10 is reported and left empty
    assert (backoff.rpe, backoff.rpe_below_5) == (None, False)
    assert [(issue.sheet, issue.cell) for issue in result.issues] == [("5w ", "C10")]
    assert "RPE '11'" in result.issues[0].message


def test_week_dates_with_mistyped_year_are_corrected(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    line = [("Day 1", [["Squat", 1, 5, None, None, 300, 7]])]
    add_old_week(sheet, 2, "Week 5", "12/23-29/24", line)
    add_old_week(sheet, 10, "Week 6", "1/6-12/24", line)
    add_old_week(sheet, 18, "Week 1", "1/6-12/25", line)
    add_old_week(sheet, 26, "Week 2", "sometime in Jan", line)
    result = parse_workbook(save(workbook, tmp_path))

    assert [row.week_start for row in result.rows] == [
        date(2024, 12, 23),
        date(2025, 1, 6),
        date(2025, 1, 6),
        None,
    ]
    assert [(issue.cell, issue.message.split(":")[0]) for issue in result.issues] == [
        ("J1", "week dates '1/6-12/24' read as 2025"),
        ("R1", "week dates '1/6-12/25' are not after the week before"),
        ("Z1", "could not read week dates 'sometime in Jan'"),
    ]


def test_loads_typed_in_kg_are_converted_per_session(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    weeks = [
        # deadlift weeks 1-3 in pounds, week 4 typed in kg (190 kg is 419 lb)
        (405, 7, 365, 6, 400, 7, 300, 7),
        (425, 7.5, 380, 6, 410, 7, 305, 7),
        (440, 8, 395, 6, 420, 7, 310, 8),
        (190, 7, 170, 6, 225, "<5", 135, 9),
    ]
    for index, (
        dl_top,
        dl_rpe,
        dl_back,
        dl_back_rpe,
        squat,
        squat_rpe,
        bench,
        bench_rpe,
    ) in enumerate(weeks):
        lines = [
            ["Deadlifts", 1, 3, None, None, dl_top, dl_rpe],
            [None, 3, 5, None, None, dl_back, dl_back_rpe],
            # the week 4 squat is a light week in pounds: as kg it would be far too heavy
            ["Squat", 1, 3, None, None, squat, squat_rpe],
            [None, 3, 5, None, None, squat, squat_rpe],
            # week 4 bench: the single could pass for kg, the set of 10 could not
            ["Bench", 1, 1, None, None, bench, bench_rpe],
            [None, 1, 10, None, None, bench, "<5"],
            # accessories are never converted
            ["DB Split Squats", 3, 8, None, None, 60 if index == 3 else 135, 8],
        ]
        add_old_week(
            sheet,
            2 + 8 * index,
            f"Week {index + 1}",
            f"6/{1 + 7 * index}-{7 + 7 * index}/25",
            [("Day 1", lines)],
        )
    result = parse_workbook(save(workbook, tmp_path))

    week_4 = [row for row in result.rows if row.week_label == "Week 4"]
    assert [(row.movement, row.assumed_kg) for row in week_4] == [
        ("Deadlifts", True),
        ("Deadlifts", True),
        ("Squat", False),
        ("Squat", False),
        ("Bench", False),
        ("Bench", False),
        ("DB Split Squats", False),
    ]
    assert week_4[0].actual_load_kg == pytest.approx(190)
    assert week_4[1].actual_load_kg == pytest.approx(170)
    assert week_4[2].actual_load_kg == lb(225)
    assert not any(row.assumed_kg for row in result.rows if row.week_label != "Week 4")
    assert [(issue.cell, "read as kg" in issue.message) for issue in result.issues] == [
        ("Z5", True),
        ("Z6", True),
    ]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("7/13-19/26", (date(2026, 7, 13), date(2026, 7, 19))),
        ("11/25-12/1/24", (date(2024, 11, 25), date(2024, 12, 1))),
        ("12/29-1/4/26", (date(2025, 12, 29), date(2026, 1, 4))),
        ("1/27 - 2/2/2025", (date(2025, 1, 27), date(2025, 2, 2))),
        ("2/30-31/25", None),
        ("next week", None),
    ],
)
def test_parse_week_dates(text, expected):
    assert _parse_week_dates(text) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, (None, False, True)),
        ("<5", (None, True, True)),
        (" < 5 ", (None, True, True)),
        (4, (None, True, True)),
        (5, (5.0, False, True)),
        (7.5, (7.5, False, True)),
        ("8.5", (8.5, False, True)),
        (10, (10.0, False, True)),
        (11, (None, False, False)),
        ("hard", (None, False, False)),
    ],
)
def test_parse_rpe(value, expected):
    assert _parse_rpe(value) == expected
