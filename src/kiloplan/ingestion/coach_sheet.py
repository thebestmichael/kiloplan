"""Parser for the coach's program spreadsheet.

Each tab is a training block. Weeks sit side by side, each led by a "Week N" cell;
days are stacked down a week, each with a header row naming the columns. Loads in
the sheet are in pounds and are converted to kg here.
"""

import re
from collections import defaultdict
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.worksheet import Worksheet

from kiloplan.analytics.e1rm import hybrid
from kiloplan.ingestion.lifts import classify

KG_PER_LB = 0.45359237
MIN_RPE = 5.0
MAX_RPE = 10.0

WEEK_LABEL = re.compile(r"^week\s*\d+", re.IGNORECASE)
DAY_LABEL = re.compile(
    r"^(day\s*\d+|monday|tuesday|wednesday|thursday|friday|saturday|sunday)$", re.IGNORECASE
)
LEADING_NUMBER = re.compile(r"^-?\d+(?:\.\d+)?")
# "7/13-19/26" or "11/25-12/1/24": start month/day, optional end month, end day, year
WEEK_DATES = re.compile(r"^(\d{1,2})/(\d{1,2})\s*-\s*(?:(\d{1,2})/)?(\d{1,2})/(\d{2,4})$")

COLUMN_LABELS = {
    "movement": "movement",
    "set": "sets",
    "sets": "sets",
    "reps": "reps",
    "intensity": "intensity",
    "targetload": "target_load",
    "trueload": "actual_load",
    "actualload": "actual_load",
    "rpelastset": "rpe",
    "lastsetrpe": "rpe",
}

# Loads are logged in pounds, but squat, bench and deadlift entries are sometimes
# typed in kg. An entry is read as kg when its e1RM as pounds is implausibly low for
# the movement in that block, and its e1RM as kg lands near the block's best.
KG_MIN_SAMPLES = 4
KG_MIN_BEST_E1RM_LB = 200.0
KG_E1RM_RANGE = (0.75, 1.10)

# A week dated this far before the one to its left has the wrong year typed in.
WRONG_YEAR_GAP = timedelta(days=300)
MAX_WEEK_GAP = timedelta(days=60)


@dataclass(frozen=True)
class SheetRow:
    """One line of the program: a movement done for `sets` x `reps`."""

    sheet: str
    cell: str
    week_index: int
    week_label: str
    week_start: date | None
    week_end: date | None
    day_label: str
    movement: str
    sets: int | None
    reps: int | None
    reps_raw: str
    intensity_raw: str
    target_load_raw: str
    actual_load_raw: str
    actual_load_kg: float | None
    rpe: float | None
    rpe_below_5: bool
    assumed_kg: bool = False


@dataclass(frozen=True)
class SheetWeek:
    """One week of a tab. `week_index` counts the tab's weeks from the left, from 0."""

    sheet: str
    week_index: int
    label: str
    start: date | None
    end: date | None


@dataclass(frozen=True)
class SheetNote:
    sheet: str
    cell: str
    week_index: int
    day_label: str
    text: str


@dataclass(frozen=True)
class Issue:
    sheet: str
    cell: str
    message: str


@dataclass
class ParseResult:
    rows: list[SheetRow] = field(default_factory=list)
    weeks: list[SheetWeek] = field(default_factory=list)
    notes: list[SheetNote] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)


def _unmangle(value: object) -> object:
    """Undo Excel's auto-conversion of numbers and ranges into dates.

    A number formatted as a date reads back as a day in 1900 (8 -> 8 Jan 1900);
    a typed range such as "8-9" reads back as that month and day of the current year.
    """
    if not isinstance(value, datetime):
        return value
    if value.year < 1905:
        return float(to_excel(value))
    return f"{value.month}-{value.day}"


def _text(value: object) -> str:
    value = _unmangle(value)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: object) -> float | None:
    """The cell as a number; text such as "60*" gives its leading number."""
    value = _unmangle(value)
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, str):
        match = LEADING_NUMBER.match(value.strip())
        if match:
            return float(match.group())
    return None


def _whole_number(value: object) -> int | None:
    """The cell as a count; ranges and text such as "8-12" or "AMRAP" give None."""
    text = _text(value)
    return int(text) if text.isdigit() else None


def _label_key(value: object) -> str:
    return re.sub(r"[^a-z]", "", _text(value).lower())


def _parse_week_dates(text: str) -> tuple[date, date] | None:
    match = WEEK_DATES.match(text.strip())
    if not match:
        return None
    start_month, start_day, end_month, end_day, year = match.groups()
    end_year = int(year) + (2000 if len(year) == 2 else 0)
    end_month = end_month or start_month
    # a week that starts in December and ends in January starts in the previous year
    start_year = end_year - 1 if int(start_month) > int(end_month) else end_year
    try:
        return (
            date(start_year, int(start_month), int(start_day)),
            date(end_year, int(end_month), int(end_day)),
        )
    except ValueError:
        return None


def _parse_rpe(value: object) -> tuple[float | None, bool, bool]:
    """Returns (rpe, below_5, valid). Valid values are "<5" and 5 to 10."""
    text = _text(value)
    if not text:
        return None, False, True
    if text.replace(" ", "") == "<5":
        return None, True, True
    number = _number(value)
    if number is None or not text.replace(".", "", 1).isdigit():
        return None, False, False
    if number < MIN_RPE:
        return None, True, True
    if number > MAX_RPE:
        return None, False, False
    return number, False, True


def _week_blocks(sheet: Worksheet) -> list[tuple[Cell, int]]:
    """Each week's label cell with the last column of its block."""
    by_row: dict[int, list[Cell]] = defaultdict(list)
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and WEEK_LABEL.match(cell.value.strip()):
                by_row[cell.row].append(cell)
    if not by_row:
        return []
    labels = max(by_row.values(), key=len)
    last_columns = [following.column - 1 for following in labels[1:]] + [sheet.max_column]
    return list(zip(labels, last_columns, strict=True))


def _find_week_dates(sheet: Worksheet, label: Cell, last_column: int) -> str:
    """The text next to the block's "Date" label, which sits just above or below the week."""
    for row in sheet.iter_rows(
        min_row=max(1, label.row - 4),
        max_row=label.row + 3,
        min_col=label.column,
        max_col=last_column,
    ):
        for index, cell in enumerate(row):
            if _label_key(cell.value) == "date":
                return next((_text(c.value) for c in row[index + 1 :] if _text(c.value)), "")
    return ""


def _next_year(day: date) -> date:
    return day.replace(year=day.year + 1, day=28 if (day.month, day.day) == (2, 29) else day.day)


def _week_dates(
    sheet: Worksheet, blocks: list[tuple[Cell, int]], issues: list[Issue]
) -> list[tuple[date | None, date | None]]:
    """Start and end date of each week, left to right, with wrong-year typos corrected."""
    dates: list[tuple[date | None, date | None]] = []
    previous: date | None = None
    for label, last_column in blocks:
        text = _find_week_dates(sheet, label, last_column)
        parsed = _parse_week_dates(text) if text else None
        if text and not parsed:
            message = f"could not read week dates {text!r}"
            issues.append(Issue(sheet.title, label.coordinate, message))
        if not parsed:
            dates.append((None, None))
            continue
        start, end = parsed
        if previous and start < previous - WRONG_YEAR_GAP:
            if previous <= _next_year(start) <= previous + MAX_WEEK_GAP:
                start, end = _next_year(start), _next_year(end)
                message = f"week dates {text!r} read as {start.year}: the year looks mistyped"
                issues.append(Issue(sheet.title, label.coordinate, message))
        if previous and start <= previous:
            message = f"week dates {text!r} are not after the week before"
            issues.append(Issue(sheet.title, label.coordinate, message))
        previous = start
        dates.append((start, end))
    return dates


def _column_value(row: tuple[Cell, ...], columns: dict[str, int], name: str) -> object:
    index = columns.get(name)
    return row[index].value if index is not None else None


def _parse_block(
    sheet: Worksheet,
    label: Cell,
    last_column: int,
    week: SheetWeek,
    result: ParseResult,
) -> None:
    columns: dict[str, int] = {}
    day_label = ""
    movement = ""
    in_notes = False

    for row in sheet.iter_rows(min_row=label.row + 1, min_col=label.column, max_col=last_column):
        first = row[0]
        first_text = _text(first.value)
        is_day = bool(DAY_LABEL.match(first_text))

        header = {
            COLUMN_LABELS[key]: index
            for index, cell in enumerate(row)
            if (key := _label_key(cell.value)) in COLUMN_LABELS
        }
        if "reps" in header:
            # older tabs have no "Movement" label: the movement shares the day column
            header.setdefault("movement", 0)
            columns = header
            day_label, movement, in_notes = first_text, "", False
            continue
        if is_day:
            day_label, movement, in_notes = first_text, "", False
        if not columns:
            continue

        movement_text = _text(_column_value(row, columns, "movement"))
        if is_day and columns["movement"] == 0:
            movement_text = ""
        if first_text.lower().startswith("notes") or movement_text.lower().startswith("notes"):
            in_notes = True
        if in_notes:
            texts = [(cell, _text(cell.value)) for cell in row]
            texts = [(cell, text) for cell, text in texts if text and _label_key(text) != "notes"]
            if texts:
                note = " ".join(text for _, text in texts)
                anchor = texts[0][0].coordinate
                result.notes.append(
                    SheetNote(sheet.title, anchor, week.week_index, day_label, note)
                )
            continue
        if movement_text:
            movement = movement_text
        if not movement or not (
            _text(_column_value(row, columns, "sets")) or _text(_column_value(row, columns, "reps"))
        ):
            continue

        anchor = row[columns["movement"]].coordinate
        rpe, rpe_below_5, rpe_valid = _parse_rpe(_column_value(row, columns, "rpe"))
        if not rpe_valid:
            message = f"RPE {_text(_column_value(row, columns, 'rpe'))!r} is not '<5' or 5 to 10"
            result.issues.append(Issue(sheet.title, anchor, message))
        load_lb = _number(_column_value(row, columns, "actual_load"))
        result.rows.append(
            SheetRow(
                sheet=sheet.title,
                cell=anchor,
                week_index=week.week_index,
                week_label=week.label,
                week_start=week.start,
                week_end=week.end,
                day_label=day_label,
                movement=movement,
                sets=_whole_number(_column_value(row, columns, "sets")),
                reps=_whole_number(_column_value(row, columns, "reps")),
                reps_raw=_text(_column_value(row, columns, "reps")),
                intensity_raw=_text(_column_value(row, columns, "intensity")),
                target_load_raw=_text(_column_value(row, columns, "target_load")),
                actual_load_raw=_text(_column_value(row, columns, "actual_load")),
                actual_load_kg=None if load_lb is None else load_lb * KG_PER_LB,
                rpe=rpe,
                rpe_below_5=rpe_below_5,
            )
        )


def _movement_key(movement: str) -> str:
    return " ".join(movement.lower().split())


def _e1rm_floor_lb(row: SheetRow) -> float | None:
    """A low estimate of the set's e1RM in pounds, reading the load as pounds.

    "<5" counts as RPE 5 and a missing RPE as RPE 10, so the estimate never overshoots.
    """
    if row.actual_load_kg is None or not row.reps:
        return None
    rpe = row.rpe if row.rpe is not None else (MIN_RPE if row.rpe_below_5 else MAX_RPE)
    return hybrid(row.actual_load_kg / KG_PER_LB, row.reps + (MAX_RPE - rpe))


def _session_key(row: SheetRow) -> tuple[str, int, str]:
    """Identifies one movement on one day of one week."""
    return _movement_key(row.movement), row.week_index, row.day_label


def _apply_kg_heuristic(rows: list[SheetRow], issues: list[Issue]) -> list[SheetRow]:
    """Re-read main-lift loads that only make sense as kg. `rows` are one sheet's rows.

    Units are typed per session, so a session is converted only when every one of its
    sets looks like kg.
    """
    estimates: dict[str, list[float]] = defaultdict(list)
    for row in rows:
        estimate = _e1rm_floor_lb(row)
        if estimate is not None:
            estimates[_movement_key(row.movement)].append(estimate)

    low, high = KG_E1RM_RANGE
    looks_like_kg: dict[tuple[str, int, str], list[bool]] = defaultdict(list)
    for row in rows:
        key = _movement_key(row.movement)
        estimate = _e1rm_floor_lb(row)
        same_movement = estimates[key]
        if estimate is None:
            continue
        looks_like_kg[_session_key(row)].append(
            classify(key).is_main
            and len(same_movement) >= KG_MIN_SAMPLES
            and (best := max(same_movement)) >= KG_MIN_BEST_E1RM_LB
            and low * best <= estimate / KG_PER_LB <= high * best
        )

    fixed = []
    for row in rows:
        votes = looks_like_kg.get(_session_key(row), [])
        if row.actual_load_kg is not None and votes and all(votes):
            best = max(estimates[_movement_key(row.movement)])
            message = (
                f"load {row.actual_load_raw} for {row.movement} read as kg: as pounds it is "
                f"far below this block's best (e1RM about {best:.0f} lb)"
            )
            issues.append(Issue(row.sheet, row.cell, message))
            row = replace(row, actual_load_kg=row.actual_load_kg / KG_PER_LB, assumed_kg=True)
        fixed.append(row)
    return fixed


def parse_workbook(source: str | Path | BinaryIO) -> ParseResult:
    """Parse every training-block tab of the coach's workbook."""
    workbook = load_workbook(source, data_only=True)
    result = ParseResult()
    for sheet in workbook.worksheets:
        blocks = _week_blocks(sheet)
        if not blocks:
            result.skipped_sheets.append(sheet.title)
            continue
        sheet_result = ParseResult()
        dates = _week_dates(sheet, blocks, sheet_result.issues)
        for index, ((label, last_column), (start, end)) in enumerate(
            zip(blocks, dates, strict=True)
        ):
            week = SheetWeek(sheet.title, index, _text(label.value), start, end)
            result.weeks.append(week)
            _parse_block(sheet, label, last_column, week, sheet_result)
        result.notes.extend(sheet_result.notes)
        result.issues.extend(sheet_result.issues)
        result.rows.extend(_apply_kg_heuristic(sheet_result.rows, result.issues))
    return result
